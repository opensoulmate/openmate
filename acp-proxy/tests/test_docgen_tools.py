"""路线图 7.2 文档生成：document 工具面（docx/pdf/xlsx）。

覆盖：docx 全块型生成回读 / xlsx 多表+表头加粗+列宽回读 / rows 简写 /
pdf 转换四态（fake converter 成功·后端不可用·转换失败·超时，docx 保留）/
参数非法显式拒绝（未知format/空content/未知块型/越界文件名）/
自动命名+落盘目录 / MEDIA: 输出 / soulmate_agent 接线护栏 + __new__ 真实
dispatch 烟测 + lazy-init 构造。（真实 LibreOffice 转换实证走"真实机器烟测"。）
"""

import subprocess
from pathlib import Path

import pytest

from agent.docgen_tools import DocGenTools, SUPPORTED_FORMATS


def _tools(tmp_path, converter=None, now=None):
    return DocGenTools(out_dir=str(tmp_path / "docs"),
                       converter=converter if converter is not None else _ok_conv,
                       now=now or (lambda: 1700000000.0))


def _ok_conv(docx_path, out_dir, soffice):
    p = Path(out_dir) / (Path(docx_path).stem + ".pdf")
    p.write_bytes(b"%PDF-1.4 fake")
    return True, "converted"


# ── docx ─────────────────────────────────────────────────────────────

def test_docx_full_blocks_roundtrip(tmp_path):
    tools = _tools(tmp_path)
    out = tools.run("docx", {
        "title": "测试报告",
        "filename": "report",
        "content": [
            {"type": "heading", "level": 1, "text": "第一章"},
            {"type": "paragraph", "text": "正文内容"},
            {"type": "bullet", "text": "要点一"},
            {"type": "table", "rows": [["名称", "数量"], ["甲", 3]]},
            {"type": "pagebreak"},
        ],
    })
    assert "✅ 文档已生成（docx）" in out
    assert "MEDIA:" in out
    path = Path(out.split("MEDIA:")[-1].strip())
    assert path.exists() and path.suffix == ".docx"
    from docx import Document
    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert "测试报告" in texts and "第一章" in texts
    assert "正文内容" in texts and "要点一" in texts
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0].cells[0].text == "名称"
    assert doc.tables[0].rows[1].cells[1].text == "3"


def test_docx_validation_errors(tmp_path):
    tools = _tools(tmp_path)
    assert "format" in tools.run("csv", {}) or "未知 format" in tools.run("csv", {})
    assert "未知 format" in tools.run("csv", {})
    assert "未知 format" in tools.run("CSV", {"content": []})  # 大小写归一
    assert "content" in tools.run("docx", {"content": []})
    assert "未知块类型" in tools.run(
        "docx", {"content": [{"type": "video", "text": "x"}]})
    assert "text 为空" in tools.run(
        "docx", {"content": [{"type": "paragraph", "text": "  "}]})
    assert "rows" in tools.run(
        "docx", {"content": [{"type": "table", "rows": []}]})
    assert "必须是对象" in tools.run("docx", {"content": ["not-a-dict"]})


# ── xlsx ─────────────────────────────────────────────────────────────

def test_xlsx_sheets_roundtrip(tmp_path):
    tools = _tools(tmp_path)
    out = tools.run("xlsx", {
        "filename": "data.xlsx",
        "sheets": [
            {"name": "销量", "rows": [["产品", "数量"], ["A", 1], ["B", 2]]},
            {"name": "备注", "rows": [["note"], ["ok"]], "header_bold": False},
        ],
    })
    assert "✅ 文档已生成（xlsx）" in out
    assert "sheets 2" in out and "rows 5" in out
    path = Path(out.split("MEDIA:")[-1].strip())
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    assert wb.sheetnames == ["销量", "备注"]
    ws = wb["销量"]
    assert ws["A1"].value == "产品" and ws["B3"].value == 2
    assert ws["A1"].font.bold is True          # 首行默认加粗
    assert wb["备注"]["A1"].font.bold is not True  # header_bold=false
    assert ws.column_dimensions["A"].width >= 8


def test_xlsx_rows_shorthand_and_errors(tmp_path):
    tools = _tools(tmp_path)
    out = tools.run("xlsx", {"rows": [["a", "b"], [1, 2]]})
    assert "✅ 文档已生成（xlsx）" in out and "rows 2" in out
    assert "sheets" in tools.run("xlsx", {})
    assert "必须是列表" in tools.run(
        "xlsx", {"sheets": [{"name": "s", "rows": ["not-a-row"]}]})
    assert "非空 rows" in tools.run(
        "xlsx", {"sheets": [{"name": "s", "rows": []}]})


# ── pdf 四态（fake converter）────────────────────────────────────────

def test_pdf_success_keeps_docx_and_reports_both(tmp_path):
    tools = _tools(tmp_path, converter=_ok_conv)
    out = tools.run("pdf", {
        "filename": "报告",
        "content": [{"type": "paragraph", "text": "中文正文"}],
    })
    assert "✅ 文档已生成（pdf）" in out
    pdf_path = Path(out.split("MEDIA:")[-1].strip())
    assert pdf_path.suffix == ".pdf" and pdf_path.exists()
    assert "同内容 docx:" in out and pdf_path.with_suffix(".docx").exists()


def test_pdf_backend_unavailable_explicit(tmp_path):
    def no_soffice(docx_path, out_dir, soffice):
        raise RuntimeError(f"libreoffice 不存在: {soffice}")

    tools = _tools(tmp_path, converter=no_soffice)
    out = tools.run("pdf", {"content": [{"type": "paragraph", "text": "x"}]})
    assert "⚠️ PDF 转换后端不可用" in out
    assert "安装方法" in out
    assert "docx 已生成" in out  # 同内容 docx 保留可交付


def test_pdf_convert_failure_explicit(tmp_path):
    def fail_conv(docx_path, out_dir, soffice):
        return False, "libreoffice 退出码 81: segfault"

    tools = _tools(tmp_path, converter=fail_conv)
    out = tools.run("pdf", {"content": [{"type": "paragraph", "text": "x"}]})
    assert "❌ PDF 转换失败" in out and "退出码 81" in out
    assert "docx 已生成" in out


def test_pdf_timeout_explicit(tmp_path):
    def slow_conv(docx_path, out_dir, soffice):
        raise subprocess.TimeoutExpired(cmd="soffice", timeout=90)

    tools = _tools(tmp_path, converter=slow_conv)
    out = tools.run("pdf", {"content": [{"type": "paragraph", "text": "x"}]})
    assert "⚠️ PDF 转换超时" in out and "未确认完成" in out


# ── 文件名/路径纪律 ──────────────────────────────────────────────────

def test_filename_sanitization(tmp_path):
    tools = _tools(tmp_path)
    for bad in ("../evil", "a/b", "a\\b", "/etc/passwd"):
        out = tools.run("docx", {
            "filename": bad, "content": [{"type": "paragraph", "text": "x"}]})
        assert "❌ 参数错误" in out, bad
    # 自动命名（无 filename）落盘在 out_dir 内
    out = tools.run("docx", {"content": [{"type": "paragraph", "text": "x"}]})
    path = Path(out.split("MEDIA:")[-1].strip())
    assert path.parent == tools.out_dir.resolve()
    assert path.name.startswith("doc-")
    # 扩展名自动补齐
    out2 = tools.run("xlsx", {"filename": "表", "rows": [["a"]]})
    assert out2.split("MEDIA:")[-1].strip().endswith(".xlsx")


# ── 接线护栏 + dispatch 烟测 ─────────────────────────────────────────

def _src() -> str:
    return (Path(__file__).resolve().parent.parent
            / "agent" / "soulmate_agent.py").read_text(encoding="utf-8")


def test_wiring_source_guards():
    src = _src()
    # 工具定义进 LLM 工具面
    assert '"name": "document"' in src
    assert "builtin_tools.extend(document_tool)" in src
    # 主循环 dispatch + code-mode 批内 dispatch 两处调用 + 共享入口定义
    assert src.count("self._document_dispatch(func_args, session_id)") == 2
    assert "def _document_dispatch" in src
    # 系统提示工具清单
    assert "- document: 文档生成" in src
    # 顶部 import
    assert "from agent.docgen_tools import DocGenTools" in src


def test_real_dispatch_smoke(tmp_path):
    from agent.soulmate_agent import SoulMateAgent
    agent = SoulMateAgent.__new__(SoulMateAgent)
    agent._docgen_tools_inst = _tools(tmp_path)
    out = agent._document_dispatch({
        "format": "docx", "filename": "smoke",
        "content": [{"type": "paragraph", "text": "hello"}]})
    assert "✅ 文档已生成（docx）" in out and "hello" not in out  # 返回统计非正文
    assert Path(out.split("MEDIA:")[-1].strip()).exists()


def test_dispatch_lazy_init_constructs_docgen(monkeypatch):
    import agent.soulmate_agent as sa

    made = {}

    class FakeTools:
        def __init__(self, *a, **k):
            made["init"] = True

        def run(self, fmt, params):
            return f"FAKE:{fmt}:{params.get('filename', '')}"

    monkeypatch.setattr(sa, "DocGenTools", FakeTools)
    agent = sa.SoulMateAgent.__new__(sa.SoulMateAgent)
    out = agent._document_dispatch({"format": "xlsx", "filename": "f"})
    assert made.get("init") is True          # lazy-init 真的构造了
    assert out == "FAKE:xlsx:f"              # 参数透传（format剥离）
