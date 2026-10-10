"""路线图 7.2 文档生成（docx/pdf/xlsx）测试。

覆盖：docx/xlsx 生成与回读（python-docx/openpyxl 真实产物）、
pdf 经 soffice 转换（真实 LibreOffice 端到端，无 soffice 自动 skip）
+ 转换失败/超时/零产出三态显式报因（注入 fake runner）、
校验面（空路径/扩展名错配/越界拒绝/空内容/形状错误/规模上限）、
原子落盘不留半截文件、MEDIA 标签、单例、soulmate_agent 接线护栏 +
__new__ 真实 dispatch 烟测。
"""

from __future__ import annotations

import os
import subprocess
from pathlib import Path

import pytest

from agent.doc_generator import (
    FORMATS,
    DocGenerator,
    DocGeneratorError,
    get_doc_generator,
    _normalize_paragraphs,
    _normalize_sheets,
    _normalize_tables,
    _resolve_target,
)

SOFFICE = pytest.mark.skipif(
    __import__("shutil").which("soffice") is None
    and __import__("shutil").which("libreoffice") is None,
    reason="无 soffice/libreoffice 二进制")


# ── docx ───────────────────────────────────────────────────────


def test_docx_generate_roundtrip(tmp_path):
    """docx 真实产物回读：标题/段落/表格单元格逐项一致 + MEDIA 标签。"""
    gen = DocGenerator()
    out = tmp_path / "报告.docx"
    text = gen.generate(
        "docx", str(out),
        title="季度报告",
        content=["第一段内容", "第二段内容"],
        tables=[{"headers": ["名称", "数量"], "rows": [["苹果", "3"], ["梨"]]}],
    )
    assert out.is_file() and out.stat().st_size > 0
    assert "✅" in text and f"MEDIA:{out}" in text and "字节" in text

    from docx import Document
    doc = Document(str(out))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "季度报告" in paras[0]
    assert "第一段内容" in paras and "第二段内容" in paras
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["名称", "数量"]
    assert [c.text for c in table.rows[1].cells] == ["苹果", "3"]
    # 短行补空不炸
    assert [c.text for c in table.rows[2].cells] == ["梨", ""]


def test_docx_content_str_form(tmp_path):
    """content 接受多行字符串（按行拆段）。"""
    gen = DocGenerator()
    out = tmp_path / "s.docx"
    text = gen.generate("docx", str(out), content="段一\n\n段二\n")
    assert "✅" in text
    from docx import Document
    paras = [p.text for p in Document(str(out)).paragraphs if p.text.strip()]
    assert paras == ["段一", "段二"]


def test_docx_extension_autocomplete(tmp_path):
    """缺扩展名自动补 .docx。"""
    gen = DocGenerator()
    out = tmp_path / "auto"
    text = gen.generate("docx", str(out), content=["x"])
    assert "✅" in text
    assert (tmp_path / "auto.docx").is_file()


# ── xlsx ───────────────────────────────────────────────────────


def test_xlsx_generate_roundtrip_typing(tmp_path):
    """xlsx 多 sheet + 数字/文本保型回读。"""
    gen = DocGenerator()
    out = tmp_path / "data.xlsx"
    text = gen.generate(
        "xlsx", str(out),
        sheets=[
            {"name": "销量", "rows": [["品名", "数量", "单价"], ["苹果", 3, 2.5]]},
            {"name": "备注", "rows": [["纯文本行"]]},
        ],
    )
    assert "✅" in text and f"MEDIA:{out}" in text

    from openpyxl import load_workbook
    wb = load_workbook(str(out))
    assert wb.sheetnames == ["销量", "备注"]
    ws = wb["销量"]
    assert [c.value for c in ws[1]] == ["品名", "数量", "单价"]
    row2 = [c.value for c in ws[2]]
    assert row2 == ["苹果", 3, 2.5]
    assert isinstance(row2[1], int) and isinstance(row2[2], float)  # 数字不是字符串


def test_xlsx_sheet_name_sanitize(tmp_path):
    """非法字符清洗/超长截断/重名去重，产物真实可打开。"""
    gen = DocGenerator()
    out = tmp_path / "names.xlsx"
    text = gen.generate(
        "xlsx", str(out),
        sheets=[
            {"name": "a[b]:c*d", "rows": [[1]]},
            {"name": "x" * 40, "rows": [[2]]},
            {"name": "dup", "rows": [[3]]},
            {"name": "dup", "rows": [[4]]},
            {"name": "", "rows": [[5]]},
        ],
    )
    assert "✅" in text
    from openpyxl import load_workbook
    names = load_workbook(str(out)).sheetnames
    assert names[0] == "a_b__c_d"
    assert len(names[1]) <= 31 and names[1].startswith("x")
    assert names[2] == "dup" and names[3] != "dup"  # 重名自动改
    assert all(n for n in names) and len(set(names)) == len(names)


# ── pdf（soffice 真实端到端 + 失败三态注入）────────────────────


@SOFFICE
def test_pdf_real_soffice_end_to_end(tmp_path):
    """真实 LibreOffice headless 转换：中文内容 → %PDF 真实产物。"""
    gen = DocGenerator()
    out = tmp_path / "文档.pdf"
    text = gen.generate(
        "pdf", str(out),
        title="中文标题",
        content=["中文段落内容，验证可选中文字。"],
    )
    assert "✅" in text and f"MEDIA:{out}" in text
    raw = out.read_bytes()
    assert raw[:5] == b"%PDF-" and out.stat().st_size > 500


def test_pdf_runner_failure_explicit(tmp_path):
    """soffice rc!=0 → 显式报因带 stderr 摘要，不产出目标文件。"""

    def runner(args, timeout):
        return subprocess.CompletedProcess(args, 1, stdout="", stderr="boom 语法错误")

    gen = DocGenerator(soffice_runner=runner)
    out = tmp_path / "f.pdf"
    text = gen.generate("pdf", str(out), content=["x"])
    assert text.startswith("[错误]") and "boom" in text
    assert not out.exists()


def test_pdf_runner_timeout_explicit(tmp_path):
    """soffice 超时 → 显式报因。"""
    def runner(args, timeout):
        raise subprocess.TimeoutExpired(cmd=args, timeout=timeout)

    gen = DocGenerator(soffice_runner=runner)
    text = gen.generate("pdf", str(tmp_path / "t.pdf"), content=["x"])
    assert text.startswith("[错误]") and "超时" in text


def test_pdf_runner_no_output_explicit(tmp_path):
    """soffice 返回成功但没产出 → 显式报因，绝不假装成功。"""
    def runner(args, timeout):
        return subprocess.CompletedProcess(args, 0, stdout="", stderr="")

    gen = DocGenerator(soffice_runner=runner)
    text = gen.generate("pdf", str(tmp_path / "n.pdf"), content=["x"])
    assert text.startswith("[错误]") and "未产出" in text


# ── 校验面 ─────────────────────────────────────────────────────


def test_reject_empty_path():
    gen = DocGenerator()
    text = gen.generate("docx", "", content=["x"])
    assert text.startswith("[错误]") and "path" in text


def test_reject_extension_mismatch(tmp_path):
    gen = DocGenerator()
    text = gen.generate("docx", str(tmp_path / "a.xlsx"), content=["x"])
    assert text.startswith("[错误]") and "不匹配" in text


def test_reject_unknown_format(tmp_path):
    gen = DocGenerator()
    text = gen.generate("pptx", str(tmp_path / "a.pptx"), content=["x"])
    assert text.startswith("[错误]") and "不支持的格式" in text
    assert set(FORMATS) == {"docx", "pdf", "xlsx"}


def test_reject_empty_content(tmp_path):
    gen = DocGenerator()
    for bad in (None, "", [], ["", "  "]):
        text = gen.generate("docx", str(tmp_path / "a.docx"), content=bad)
        assert text.startswith("[错误]") and "content" in text


def test_reject_content_wrong_type(tmp_path):
    gen = DocGenerator()
    text = gen.generate("docx", str(tmp_path / "a.docx"), content={"a": 1})
    assert text.startswith("[错误]") and "形状" in text


def test_reject_paragraph_limit(tmp_path, monkeypatch):
    import agent.doc_generator as dg
    monkeypatch.setattr(dg, "MAX_PARAGRAPHS", 2)
    text = DocGenerator().generate("docx", str(tmp_path / "a.docx"),
                                  content=["1", "2", "3"])
    assert text.startswith("[错误]") and "超限" in text


def test_reject_bad_tables_shapes(tmp_path):
    gen = DocGenerator()
    cases = [
        "not-a-list",
        [{"rows": [[1]]}],                       # headers 缺
        [{"headers": ["a"], "rows": "x"}],       # rows 非列表
        [{"headers": ["a"], "rows": [["x"]]}],   # 合法对照放在合法用例
        [{"headers": ["a"], "rows": ["x"]}],     # 行非列表
        "table-dict",                            # 非 dict 条目
    ]
    text = gen.generate("docx", str(tmp_path / "a.docx"),
                        content=["x"], tables=cases[0])
    assert text.startswith("[错误]")
    for bad in cases[1:3] + cases[4:]:
        text = gen.generate("docx", str(tmp_path / "a.docx"),
                            content=["x"], tables=[bad] if isinstance(bad, dict) else bad)
        assert text.startswith("[错误]"), bad
    # 合法表格正常产出
    ok = gen.generate("docx", str(tmp_path / "b.docx"), content=["x"],
                      tables=[{"headers": ["a"], "rows": [["x"]]}])
    assert "✅" in ok


def test_reject_bad_sheets(tmp_path):
    gen = DocGenerator()
    assert gen.generate("xlsx", str(tmp_path / "a.xlsx"),
                        sheets=[]).startswith("[错误]")
    assert gen.generate("xlsx", str(tmp_path / "a.xlsx"),
                        sheets=[{"name": "s"}]).startswith("[错误]")  # rows 空
    assert gen.generate("xlsx", str(tmp_path / "a.xlsx"),
                        sheets=[{"name": "s", "rows": [["a"], "bad"]}]).startswith("[错误]")
    assert gen.generate("xlsx", str(tmp_path / "a.xlsx"),
                        sheets=["bad"]).startswith("[错误]")


def test_reject_path_outside_roots(tmp_path):
    """默认只允许 $HOME 与 /tmp 下；越界拒绝且不落任何文件。"""
    gen = DocGenerator()
    text = gen.generate("docx", "/etc/soulmate_越界测试.docx", content=["x"])
    assert text.startswith("[错误]") and "越界" in text
    assert not Path("/etc/soulmate_越界测试.docx").exists()


def test_env_roots_override(tmp_path, monkeypatch):
    """SOULMATE_DOC_ROOTS 自定义根目录生效。"""
    root = tmp_path / "允许根"
    root.mkdir()
    monkeypatch.setenv("SOULMATE_DOC_ROOTS", str(root))
    gen = DocGenerator()
    assert "✅" in gen.generate("docx", str(root / "in.docx"), content=["x"])
    assert gen.generate("docx", str(tmp_path / "out.docx"),
                        content=["x"]).startswith("[错误]")  # 原 tmp 根不再允许


def test_normalize_helpers_direct():
    """归一函数显式拒绝（直接调用面）。"""
    with pytest.raises(DocGeneratorError):
        _resolve_target("", "docx")
    with pytest.raises(DocGeneratorError):
        _resolve_target("/tmp/x", "badfmt")
    with pytest.raises(DocGeneratorError):
        _normalize_paragraphs(123)
    with pytest.raises(DocGeneratorError):
        _normalize_tables([{"headers": [], "rows": []}])
    with pytest.raises(DocGeneratorError):
        _normalize_sheets(None)
    assert _normalize_tables(None) == []


# ── 原子落盘 / 单例 ────────────────────────────────────────────


def test_no_half_file_on_save_failure(tmp_path, monkeypatch):
    """save 抛异常：目标文件不存在、tmp 不残留。"""
    import agent.doc_generator as dg

    def boom(tmp):
        tmp.write_bytes(b"half")
        raise RuntimeError("写入炸了")

    target = tmp_path / "x.docx"
    with pytest.raises(RuntimeError):
        dg._atomic_save(boom, target)
    assert not target.exists()
    assert not list(tmp_path.glob(".*tmp*"))  # 无半截残留


def test_generate_error_text_on_unexpected_exception(tmp_path, monkeypatch):
    """生成内部抛非 DocGeneratorError 也文本化，不抛穿。"""
    gen = DocGenerator()
    monkeypatch.setattr(gen, "_write_docx",
                        lambda *a, **k: (_ for _ in ()).throw(OSError("磁盘炸了")))
    text = gen.generate("docx", str(tmp_path / "a.docx"), content=["x"])
    assert text.startswith("[错误]") and "磁盘炸了" in text


def test_singleton():
    assert get_doc_generator() is get_doc_generator()
    assert isinstance(get_doc_generator(), DocGenerator)


# ── soulmate_agent 接线护栏 + 真实 dispatch 烟测 ────────────────


def test_soulmate_agent_wiring_guards():
    src = Path(__file__).resolve().parent.parent.joinpath(
        "agent/soulmate_agent.py").read_text(encoding="utf-8")
    # 工具定义
    assert '"name": "doc_generate"' in src
    # 两处dispatch点：主循环 + code-mode内层
    assert src.count('func_name == "doc_generate"') >= 2
    # 分发helper + lazy-init + 顶部import
    assert "def _doc_dispatch" in src
    assert "_get_doc_generator" in src
    assert "from agent.doc_generator import" in src
    # 系统提示工具清单（f-string 内 {{ }} 转义后仍能命中关键词）
    assert "doc_generate: 文档生成" in src


def test_doc_dispatch_smoke(monkeypatch, tmp_path):
    """__new__构造+lazy-init真实dispatch烟测：xlsx/docx/错误三态。"""
    import agent.soulmate_agent as sa

    gen = DocGenerator()  # 真实 docx/xlsx 产物（不碰 soffice）
    monkeypatch.setattr(sa, "_get_doc_generator", lambda: gen)
    agent = sa.SoulMateAgent.__new__(sa.SoulMateAgent)  # 不过__init__

    out = agent._doc_dispatch({
        "format": "xlsx", "path": str(tmp_path / "烟测.xlsx"),
        "sheets": [{"name": "s", "rows": [["a", 1]]}],
    })
    assert "✅" in out and "MEDIA:" in out and (tmp_path / "烟测.xlsx").is_file()

    out = agent._doc_dispatch({
        "format": "docx", "path": str(tmp_path / "烟测.docx"),
        "title": "T", "content": ["c"],
    })
    assert "✅" in out and (tmp_path / "烟测.docx").is_file()

    out = agent._doc_dispatch({"format": "docx", "path": "", "content": ["c"]})
    assert out.startswith("[错误]")
    # lazy-init 已缓存实例
    assert getattr(agent, "_doc_generator", None) is gen


def test_doc_dispatch_prompt_fstring_builds():
    """系统提示是 f-string：doc_generate 行的 {{ }} 转义正确，构造不炸。"""
    import agent.soulmate_agent as sa
    agent = sa.SoulMateAgent.__new__(sa.SoulMateAgent)
    prompt = agent._build_base_system_prompt("/tmp")
    assert "doc_generate: 文档生成" in prompt
    assert "{{" not in prompt  # 转义被正确还原
